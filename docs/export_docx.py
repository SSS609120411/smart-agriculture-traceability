# -*- coding: utf-8 -*-
"""将 docs/ 目录下的 Markdown 文档批量导出为 Word（.docx）。

用法（在 docs 目录下）：
    python export_docx.py
"""
import os
import re

from docx import Document
from docx.shared import Pt

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_runs(paragraph, text):
    """解析行内 **粗体** 与 `代码` 并写入段落。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def convert(md_path, docx_path):
    doc = Document()
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # 代码块
        if line.strip().startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过结尾 ```
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue

        # 表格
        if line.strip().startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            rows = [r for r in rows if not re.match(r"^\s*\|[\s:\-|]+\|\s*$", r)]
            if rows:
                header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
                table = doc.add_table(rows=len(rows), cols=len(header))
                table.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    cells = [c.strip() for c in row.strip().strip("|").split("|")]
                    for ci, cell in enumerate(cells):
                        if ci < len(header):
                            add_runs(table.rows[ri].cells[ci].paragraphs[0], cell)
                            if ri == 0:
                                for run in table.rows[ri].cells[ci].paragraphs[0].runs:
                                    run.bold = True
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2), level=level)
            i += 1
            continue

        # 列表
        if line.strip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line.strip()[2:])
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, line.strip())
        i += 1

    doc.save(docx_path)
    print(f"  生成 {os.path.basename(docx_path)}")


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    files = sorted(glob for glob in [f for f in os.listdir(here) if f.endswith(".md")])
    for f in files:
        convert(os.path.join(here, f), os.path.join(here, f[:-3] + ".docx"))
    print("全部完成！")


if __name__ == "__main__":
    main()
